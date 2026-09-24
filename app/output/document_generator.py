"""Create the final .docx answer document.

WHAT: Turns an assignment + its AI solution into a polished Word document
      and saves it to data/generated/.
WHY:  The submission-and-review loop needs a real document the student can
      read, edit and attach — not just terminal output.
HOW:  python-docx builds the file: title cover info (course, title, date,
      due date, IDs), a "Question / Task" section, then the AI solution.
      The solution is written in a light Markdown subset so lines like
      '## Final Answer' become real Word headings and '- ' becomes bullets.
"""

from datetime import date
from pathlib import Path

from docx import Document


def generate_docx(assignment_info: dict, solution_text: str, out_path) -> str:
    """Write the answer document. Returns the path it was saved to."""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # ---- Cover block -------------------------------------------------
    doc.add_heading("Classroom AI Assistant — Solution", 0)
    doc.add_heading(assignment_info.get("assignment_title", "Untitled assignment"), level=1)

    meta = [
        ("Course", assignment_info.get("course", "")),
        ("Generated", date.today().isoformat()),
        ("Due date", assignment_info.get("due_date") or "(no due date)"),
        ("Assignment ID", assignment_info.get("assignment_id", "")),
        ("Course ID", assignment_info.get("course_id", "")),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(str(value))

    # ---- Question / task ---------------------------------------------
    doc.add_heading("Question / Task", level=1)
    description = assignment_info.get("description", "").strip()
    doc.add_paragraph(description if description else "(No description was provided by the teacher.)")

    # ---- Solution ----------------------------------------------------
    doc.add_page_break()
    doc.add_heading("Solution", level=1)
    _write_solution(doc, solution_text)

    doc.save(out_path)
    return str(out_path)


def _write_solution(doc: Document, text: str) -> None:
    """Render the AI solution with simple Markdown as styled Word content."""
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style="List Bullet")
        elif _looks_like_numbered(stripped):
            doc.add_paragraph(stripped, style="List Number")
        else:
            doc.add_paragraph(stripped)

        if _is_heading(stripped):
            continue


def _is_heading(stripped: str) -> bool:
    return stripped.startswith("#")


def _looks_like_numbered(line: str) -> bool:
    """Heuristic: lines like '1.', '1)', '2.' as numbered list starts."""
    if line[:2].replace(".", "").isdigit() and line[:2].strip():
        return True
    return len(line) > 2 and line[0].isdigit() and line[1] in ".):"